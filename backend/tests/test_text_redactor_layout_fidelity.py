"""文本型 PDF 替换成品的排版保真（Issue #84）：

根因②修复——实体文本带 PyMuPDF 空格伪影（如 `2018 年3 月26 日`），与
pdf2docx 产出 docx 的空格位置不一致（尤其跨 run 边界），精确匹配漏替
→ 残留校验拦截 → 回退原位替换。docx 替换匹配须做空白归一。

根因①表现层修复——原位替换（兜底）的替换标签字号不再钳 10pt 帽，
按原文字 span 的真实字号写入，超宽自适应缩放。
"""

import asyncio

import fitz
import pytest

from app.models.common import ReplacementMode
from app.models.entity_schemas import Entity
from app.services.redaction.replacement_strategy import RedactionContext
from app.services.redaction.text_redactor import TextRedactorMixin


def _ctx(replacements: dict[str, str]) -> RedactionContext:
    ctx = RedactionContext(mode=ReplacementMode.CUSTOM)
    ctx.set_custom_replacements(replacements)
    return ctx


def _ent(text: str, type_: str = "DATE") -> Entity:
    return Entity(
        id=f"e-{text}", text=text, type=type_, start=0, end=len(text),
        page=1, selected=True,
    )


class _Para:
    """最小段落桩：N 个 run。"""

    class _R:
        def __init__(self, text):
            self.text = text

    def __init__(self, *runs):
        self.runs = [self._R(t) for t in runs]


# ---------------------------------------------------------- 修复② 空白归一匹配


def test_run_level_replace_space_artifact_across_runs():
    """实体空格位置与 docx 段落不一致（跨 run 边界空格丢失）也能替换。

    真实案卷形态（#84 根因②）：实体 `2018 年3 月26 日`，docx 两 run
    拼接后是 `自2018 年3` + `月26 日至…` = `2018 年3月26 日`。
    """
    r = TextRedactorMixin()
    para = _Para("罪被法院判处拘役三个月（自2018 年3", "月26 日止），并处罚金二千元")
    count = r._replace_in_paragraph(para, {"2018 年3 月26 日": "【日期1】"})
    assert count == 1
    full = "".join(run.text for run in para.runs)
    assert "【日期1】" in full
    assert "2018 年3" not in full


def test_run_level_replace_entity_spaces_nowhere_in_docx():
    """实体空格伪影与 docx 完全不同（`2021 年01 月01 日` 0 命中形态）。"""
    r = TextRedactorMixin()
    para = _Para("自2021年01月01日起至2022年01月01日止")
    count = r._replace_in_paragraph(para, {"2021 年01 月01 日": "【日期1】"})
    assert count == 1
    assert para.runs[0].text == "自【日期1】起至2022年01月01日止"


def test_run_level_replace_without_spaces_unchanged_behavior():
    """无空格实体的既有行为不回退（含跨 run 拆分）。"""
    r = TextRedactorMixin()
    para = _Para("犯罪嫌疑人", "张三丰到庭")
    count = r._replace_in_paragraph(para, {"张三丰": "【人1】"})
    assert count == 1
    assert para.runs[0].text == "犯罪嫌疑人"
    assert para.runs[1].text == "【人1】到庭"


def test_run_level_longest_match_wins_with_squeeze():
    """空白归一后仍保持最长匹配优先（张三丰 不被 张三 提前吞掉）。"""
    r = TextRedactorMixin()
    para = _Para("张三丰到庭")
    count = r._replace_in_paragraph(para, {"张三": "【人1】", "张三丰": "【人2】"})
    assert count == 1
    assert para.runs[0].text == "【人2】到庭"


def test_xml_level_replace_with_space_artifact_mismatch():
    """XML 趟（嵌套/混合段落）同样空白归一匹配。"""
    from lxml import etree

    NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xml = (
        f'<w:p xmlns:w="{NS}">'
        f'<w:r><w:t>罪被法院判处拘役三个月（自2018 年3</w:t></w:r>'
        f'<w:r><w:t>月26 日止），并处罚金</w:t></w:r>'
        f"</w:p>"
    )
    p = etree.fromstring(xml)
    r = TextRedactorMixin()
    count = r._replace_in_docx_xml_paragraph(p, {"2018 年3 月26 日": "【日期1】"})
    assert count == 1
    full = "".join(t.text or "" for t in r._docx_xpath(p, ".//w:t"))
    assert "【日期1】" in full


# ---------------------------------------------------------- 修复① 原字号替换


def _make_pdf(path, text, fontsize=16):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), text, fontsize=fontsize, fontname="china-s")
    doc.save(str(path))
    doc.close()


@pytest.fixture()
def _allow_tmp_upload(tmp_path, monkeypatch):
    """open_pdf_checked 只允许 UPLOAD_DIR/OUTPUT_DIR 内的路径。"""
    from app.core.config import settings

    up = tmp_path / "uploads"
    up.mkdir()
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(up))
    # 源文件放进白名单目录
    return up


def test_pdf_inplace_replacement_uses_original_font_size(tmp_path, _allow_tmp_upload):
    """原位替换：原文字 16pt，替换文本也应 ≈16pt（不再钳 10pt 帽）。"""
    src = _allow_tmp_upload / "in.pdf"
    out = tmp_path / "out.pdf"
    _make_pdf(src, "张三与李四签订协议。", 16)
    r = TextRedactorMixin()
    asyncio.run(
        r._redact_pdf_text(
            str(src), str(out),
            [_ent("张三", "PERSON"), _ent("李四", "PERSON")],
            _ctx({"张三": "人甲", "李四": "人乙"}),
        )
    )
    out_doc = fitz.open(str(out))
    spans = [
        s
        for b in out_doc[0].get_text("dict")["blocks"]
        for l in b.get("lines", [])
        for s in l["spans"]
    ]
    out_doc.close()
    repl = [s for s in spans if "人甲" in s["text"] or "人乙" in s["text"]]
    assert repl, f"应能提取到替换文本 spans: {[s['text'] for s in spans]}"
    for s in repl:
        assert s["size"] > 12, f"替换字号应接近原文 16pt，实际 {s['size']}"


def test_pdf_inplace_replacement_shrinks_when_longer(tmp_path, _allow_tmp_upload):
    """替换词远长于原文时宽度自适应缩放，且不整体丢失。"""
    src = _allow_tmp_upload / "in.pdf"
    out = tmp_path / "out.pdf"
    _make_pdf(src, "张三", 16)
    r = TextRedactorMixin()
    asyncio.run(
        r._redact_pdf_text(
            str(src), str(out), [_ent("张三", "PERSON")],
            _ctx({"张三": "某省某市某区人民法院刑事审判庭某法官"} ),
        )
    )
    out_doc = fitz.open(str(out))
    text = out_doc[0].get_text()
    out_doc.close()
    assert "某省某市" in text, f"长替换词不应整体丢失，实际: {text!r}"


# ---------------------------------------------------------- 端到端（docx 主路）


def test_docx_path_no_residual_with_space_artifact(tmp_path):
    """docx 主路：空格错位实体（跨 run）替换后无残留，不再触发回退。"""
    from docx import Document as Docx

    docx_path = tmp_path / "source.docx"
    out = tmp_path / "out.docx"
    d = Docx()
    p = d.add_paragraph()
    p.add_run("自2018 年3")
    p.add_run("月26 日起至2019 年01 月01 日止。")
    d.save(str(docx_path))

    r = TextRedactorMixin()
    ents = [_ent("2018 年3 月26 日"), _ent("2019 年01 月01 日")]
    count = asyncio.run(
        r._redact_docx(
            str(docx_path), str(out), ents,
            _ctx({"2018 年3 月26 日": "【日期1】", "2019 年01 月01 日": "【日期2】"}),
        )
    )
    assert count == 2, "两个空格错位实体都应被替换"
    residual = r._docx_texts_present(str(out), {e.text for e in ents})
    assert not residual, f"替换后不应残留原文: {residual}"
