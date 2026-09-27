"""独立审计复现测试：超链接/修订插入(w:ins)/域代码(instrText)内的实体
不因 pass1 整段标记已处理而被 pass2 跳过；词池导入裸字典与字符串 words 防护。"""

import asyncio
import zipfile
from io import BytesIO

import pytest
from docx import Document
from lxml import etree as et

from app.core.config import settings
from app.models.entity_schemas import Entity
from app.models.redaction_schemas import RedactionConfig
from app.services.redactor import Redactor

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@pytest.fixture()
def _dirs(tmp_path, monkeypatch):
    up = tmp_path / "uploads"; out = tmp_path / "outputs"
    up.mkdir(); out.mkdir()
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(up))
    monkeypatch.setattr(settings, "OUTPUT_DIR", str(out))
    return up, out


def _make_docx(path):
    """正文段落含：超链接内 run、修订插入 w:ins 内 run、域代码 instrText。"""
    doc = Document()
    doc.add_paragraph("正文：马国栋")
    buf = BytesIO(); doc.save(buf)
    zin = zipfile.ZipFile(BytesIO(buf.getvalue()))
    root = et.fromstring(zin.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    p = et.SubElement(body, f"{{{W}}}p")

    r1 = et.SubElement(p, f"{{{W}}}r")
    et.SubElement(r1, f"{{{W}}}t").text = "超链接："
    hl = et.SubElement(p, f"{{{W}}}hyperlink")
    hr = et.SubElement(hl, f"{{{W}}}r")
    et.SubElement(hr, f"{{{W}}}t").text = "陈文清的个人主页"

    p2 = et.SubElement(body, f"{{{W}}}p")
    ins = et.SubElement(p2, f"{{{W}}}ins")
    ir = et.SubElement(ins, f"{{{W}}}r")
    et.SubElement(ir, f"{{{W}}}t").text = "修订插入：范立群"

    p3 = et.SubElement(body, f"{{{W}}}p")
    fr = et.SubElement(p3, f"{{{W}}}r")
    et.SubElement(fr, f"{{{W}}}instrText").text = " REF 吴永康 \\h "

    out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in zin.infolist():
            data = (et.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    if item.filename == "word/document.xml" else zin.read(item.filename))
            z.writestr(item, data)
    with open(path, "wb") as f:
        f.write(out.getvalue())


@pytest.mark.asyncio
async def test_hyperlink_ins_instrtext_redacted(_dirs):
    up, _ = _dirs
    src = up / "t.docx"
    _make_docx(str(src))
    entities = [
        Entity(id="e1", text="马国栋", type="PERSON", start=0, end=3, page=1, selected=True),
        Entity(id="e2", text="陈文清", type="PERSON", start=0, end=3, page=1, selected=True),
        Entity(id="e3", text="范立群", type="PERSON", start=0, end=3, page=1, selected=True),
        Entity(id="e4", text="吴永康", type="PERSON", start=0, end=3, page=1, selected=True),
    ]
    pools = {"PERSON": {"words": ["甲某", "乙某", "丙某", "丁某"], "strategy": "numbered", "custom_map": {}}}
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "docx"},
        entities=entities, bounding_boxes=[],
        config=RedactionConfig(replacement_mode="pseudonym", word_pools=pools),
    )
    with zipfile.ZipFile(result["output_path"]) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    for gone in ["马国栋", "陈文清", "范立群", "吴永康"]:
        assert gone not in xml, f"嵌套内容实体残留: {gone}"
    for present in ["甲某", "乙某", "丙某", "丁某"]:
        assert present in xml, f"替换词缺失: {present}"
    assert result["residual_entities"] == []


def test_import_bare_pool_dict(tmp_path, monkeypatch):
    from app.core.config import settings as cfg
    from app.services import word_pool_service as svc

    monkeypatch.setattr(cfg, "WORD_POOL_STORE_PATH", str(tmp_path / "wp.json"))
    count = svc.import_word_pools({"PERSON": {"words": ["赵大"], "strategy": "cycle", "custom_map": {}}},
                                  owner_id="u9")
    assert count == 1
    merged = svc.load_word_pools(owner_id="u9")
    assert merged["PERSON"]["words"] == ["赵大"]


def test_normalize_pool_string_words():
    from app.services.word_pool_service import _normalize_pool

    pool = _normalize_pool({"words": "张三李四", "strategy": "numbered"})
    assert pool["words"] == [], "字符串 words 应被丢弃而非逐字拆分"


@pytest.mark.asyncio
async def test_entity_spanning_run_and_hyperlink_redacted(_dirs):
    """实体横跨「直接 run ↔ 超链接」边界：run 级与嵌套级单独都拼不出
    完整原文，必须由整段 XML 趟兜住；自检也不得误报干净。"""
    up, _ = _dirs
    src = up / "span.docx"
    doc = Document()
    doc.add_paragraph("另一处完整出现：马国栋")
    buf = BytesIO(); doc.save(buf)
    zin = zipfile.ZipFile(BytesIO(buf.getvalue()))
    root = et.fromstring(zin.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    p = et.SubElement(body, f"{{{W}}}p")
    for text in ("客户", "马"):
        r = et.SubElement(p, f"{{{W}}}r")
        et.SubElement(r, f"{{{W}}}t").text = text
    hl = et.SubElement(p, f"{{{W}}}hyperlink")
    hr = et.SubElement(hl, f"{{{W}}}r")
    et.SubElement(hr, f"{{{W}}}t").text = "国栋"
    r2 = et.SubElement(p, f"{{{W}}}r")
    et.SubElement(r2, f"{{{W}}}t").text = "签收。"
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in zin.infolist():
            data = (et.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    if item.filename == "word/document.xml" else zin.read(item.filename))
            z.writestr(item, data)
    with open(src, "wb") as f:
        f.write(out.getvalue())

    entities = [Entity(id="e1", text="马国栋", type="PERSON", start=0, end=3, page=1, selected=True)]
    pools = {"PERSON": {"words": ["甲某"], "strategy": "numbered", "custom_map": {}}}
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "docx"},
        entities=entities, bounding_boxes=[],
        config=RedactionConfig(replacement_mode="pseudonym", word_pools=pools),
    )
    with zipfile.ZipFile(result["output_path"]) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "马国栋" not in xml and "马" + "国栋" not in xml.replace("</w:t>", "").replace('<w:t xml:space="preserve">', ""), \
        "跨界实体残留"
    # 直接 run 里的「马」与超链接里的「国栋」拼成原文也算残留——XML 逐节点断言
    assert "国栋" not in xml, "超链接内半个实体残留"
    assert "甲某" in xml
    assert result["residual_entities"] == []


@pytest.mark.asyncio
async def test_mixed_paragraph_cross_key_no_wrong_pseudonym(_dirs):
    """混合段落（run 键 + 跨界键并存）：张三→李四、李四→王五 且李四跨界。
    run/union 分流若两趟先后跑，union 趟会把 run 趟刚写入的「李四」再改成
    王五（错误化名）；混合段落必须单趟完成。"""
    up, _ = _dirs
    src = up / "mixed.docx"
    doc = Document()
    doc.add_paragraph("无关段落")
    buf = BytesIO(); doc.save(buf)
    zin = zipfile.ZipFile(BytesIO(buf.getvalue()))
    root = et.fromstring(zin.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    p = et.SubElement(body, f"{{{W}}}p")
    r = et.SubElement(p, f"{{{W}}}r")
    et.SubElement(r, f"{{{W}}}t").text = "张三与李"
    hl = et.SubElement(p, f"{{{W}}}hyperlink")
    hr = et.SubElement(hl, f"{{{W}}}r")
    et.SubElement(hr, f"{{{W}}}t").text = "四同行"
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in zin.infolist():
            data = (et.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    if item.filename == "word/document.xml" else zin.read(item.filename))
            z.writestr(item, data)
    with open(src, "wb") as f:
        f.write(out.getvalue())

    entities = [
        Entity(id="e1", text="张三", type="PERSON", start=0, end=2, page=1, selected=True),
        Entity(id="e2", text="李四", type="PERSON", start=0, end=2, page=1, selected=True),
    ]
    cfg = RedactionConfig(
        replacement_mode="custom",
        custom_replacements={"张三": "李四", "李四": "王五"},
    )
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "docx"},
        entities=entities, bounding_boxes=[], config=cfg,
    )
    with zipfile.ZipFile(result["output_path"]) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "张三" not in xml
    # 正确结果：张三→李四、李四→王五（不是全部变王五）
    assert xml.count("王五") == 1 and xml.count("李四") == 1, \
        f"跨键串词: 李四出现{xml.count('李四')}次 王五出现{xml.count('王五')}次"
    # 李四作为张三的替换词出现在成品中，自检按词池回灌语义如实告警
    assert result["residual_entities"] == ["李四"]
