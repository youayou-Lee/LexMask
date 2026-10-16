"""零泄漏闸（AC3 红线）：对照表强归一化 grep + 可选成品复扫。

语义对齐旧 eval/scripts/leak_check.py（旧文件冻结只读）：
_norm 去空白/标点/连字符 + 大写折叠（强于打分口径的 _normalize）；
对照表 CSV 列 = 原文,类型,化名（utf-8-sig）。
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from ..protocol import Answer, AnswerForm, ExamItem
from .base import GateResult

_PUNCT = re.compile(r"[\s\-—_·・.,，。;；:：!！?？'‘’\"“”()（）\[\]【】<>《》+*=|\\/]+")


def norm(value: str) -> str:
    """强归一化：与 leak_check._norm 同语义（测试断言两实现同判）。"""
    return _PUNCT.sub("", value).upper()


def load_mapping(mapping_csv: Path) -> list[dict]:
    """CSV 列：原文,类型,化名。缺列报错（同 leak_check.load_mapping 语义）。"""
    rows = []
    with open(mapping_csv, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        for col in ("原文", "类型", "化名"):
            if col not in (reader.fieldnames or []):
                raise ValueError(f"对照表缺列 {col!r}: {mapping_csv}")
        for row in reader:
            if row.get("原文"):
                rows.append({"原文": row["原文"], "类型": row["类型"], "化名": row["化名"]})
    return rows


def _extract_text(artifact_path: Path) -> str:
    """成品文本抽取：txt/md/html/rtf 直读；pdf 用 fitz；docx 用 python-docx（含表格）。

    与 leak_check.extract_text 同语义；依赖缺失或解析失败抛错由 runner 记 FAIL。
    """
    suffix = artifact_path.suffix.lower()
    if suffix in (".txt", ".md", ".html", ".rtf"):
        return artifact_path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        import fitz

        with fitz.open(artifact_path) as doc:
            return "\n".join(page.get_text() for page in doc)
    if suffix == ".docx":
        import docx

        d = docx.Document(str(artifact_path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    raise ValueError(f"零泄漏闸不支持成品格式: {suffix}（{artifact_path}）")


def gate(answer: Answer, item: ExamItem, env: dict) -> GateResult:
    if answer.form is not AnswerForm.ARTIFACT_FILE:
        return GateResult("leak_free", True, {"skipped": f"form={answer.form.value}"})
    sidecars = (item.payload or {}).get("sidecars") or {}
    mapping_path = sidecars.get("mapping_csv")
    if not mapping_path or not answer.artifact_path:
        return GateResult("leak_free", False, {"reason": "缺对照表 sidecar 或成品文件"})
    try:
        rows = load_mapping(Path(mapping_path))
        text = _extract_text(Path(answer.artifact_path))
    except Exception as exc:  # noqa: BLE001 — 闸门侧任何失败=不通过
        return GateResult("leak_free", False, {"reason": f"执行失败: {exc}"})
    norm_text = norm(text)
    findings = []
    for row in rows:
        if row["原文"] and norm(row["原文"]) in norm_text:
            findings.append({"原文": row["原文"], "类型": row["类型"], "how": "normalized-grep"})
    passed = not findings
    details: dict = {
        "checked_originals": len(rows),
        "findings": findings,
        "rescan_plan": None,
    }
    # 步骤②成品复扫：需要 api 句柄（非 --ci、真服务环境）才跑
    if env.get("api_handles"):
        details["rescan_plan"] = {"not_implemented": True, "skipped": "复扫在实例验收阶段由背靠背脚本调用 run_rescan 语义执行"}
    return GateResult("leak_free", passed, details)
